from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import HotSearchBatch, HotSearchConfig, HotSearchTerm, Product, Rule


HEADER_ALIASES = {
    "rank": ("排名", "排行", "序号"),
    "category": ("搜索词主分类", "主分类", "类目", "三级分类"),
    "keyword": ("关键词", "搜索词", "热搜词"),
    "search_uv_index": ("搜索UV指数", "搜索UV指数 ", "UV指数"),
    "search_uv_growth": ("搜索UV指数涨幅", "搜索UV指数涨幅 ", "UV涨幅"),
    "click_rate_index": ("搜索点击率指数", "搜索点击率指数 ", "点击率指数"),
    "opportunity_index": ("机会指数", "机会指数 "),
    "gmv_index": ("成交金额指数", "成交金额指数 "),
    "sales_index": ("销售量指数", "销售量指数 "),
}


DEFAULT_AVOID_TERMS = {
    "中幼童": "婴,旗舰店,戴维贝拉,北面,波司登,左西,蕉下,babycare,gap,巴拉巴拉,jnbybyjnby,嘟嘟家,kk树,jk",
    "中童": "婴,宝宝,旗舰店,戴维贝拉,北面,波司登,左西,蕉下,babycare,361,斐乐,江南布衣,gap,巴拉巴拉,jnbybyjnby,安德玛,嘟嘟家,kk树,jk",
    "婴幼童": "大童,童泰,英氏,babylove,贝肽斯,domiamia,戴维贝拉,旗舰店,北面,波司登,左西,蕉下,babycare,gap,斐乐,巴拉巴拉,jnbybyjnby,嘟嘟家,kk树,jk",
    "鞋品": "江博士,泰兰尼斯,基诺浦,卡特兔,牧童,巴布豆,VANS,kappa,李宁,旗舰店,戴维贝拉,波司登,安踏,斯凯奇,阿迪达斯,abckids,巴拉巴拉,耐克,361,kk树,jk",
    "HOME": "丽婴房,英氏,旗舰店,戴维贝拉,波司登,爱慕,likeuu,蕉内,361,斐乐,巴拉巴拉,嘟嘟家,kk树,jk",
    "用品": "北面,巴拉巴拉,kk树,jk",
    "婴童": "大童,童泰,英氏,babylove,贝肽斯,domiamia,戴维贝拉,旗舰店,北面,波司登,左西,babycare,gap,巴拉巴拉,jnbybyjnby,kk树,jk",
}


def _clean(value) -> str:
    return str(value or "").strip()


def _number(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _int(value) -> int | None:
    number = _number(value)
    return int(number) if number is not None else None


def _headers(row: tuple) -> dict[str, int]:
    raw = {_clean(cell): index for index, cell in enumerate(row)}
    mapping: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if alias in raw:
                mapping[field] = raw[alias]
                break
    return mapping


def _value(row: tuple, mapping: dict[str, int], field: str):
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def _split_terms(content: str | None) -> list[str]:
    return [term.strip() for term in re.split(r"[,，、\n]+", content or "") if term.strip()]


def _product_line(product: Product) -> str:
    return _clean(product.age_range)


def seed_default_hot_search_avoid_rules(db: Session) -> None:
    existing = db.query(Rule).filter(Rule.rule_type == "hot_search_avoid_terms").first()
    if existing:
        return
    for line, terms in DEFAULT_AVOID_TERMS.items():
        db.add(
            Rule(
                rule_type="hot_search_avoid_terms",
                rule_name=line,
                content=terms,
                enabled=True,
                updated_by="system",
            )
        )
    db.commit()


def get_hot_search_config(db: Session) -> dict:
    config = db.query(HotSearchConfig).order_by(HotSearchConfig.id.asc()).first()
    return {
        "enabled_by_default": bool(config.enabled_by_default) if config else False,
        "updated_by": config.updated_by if config else None,
    }


def set_hot_search_config(db: Session, enabled_by_default: bool, updated_by: str = "operator") -> dict:
    config = db.query(HotSearchConfig).order_by(HotSearchConfig.id.asc()).first()
    if not config:
        config = HotSearchConfig()
        db.add(config)
    config.enabled_by_default = enabled_by_default
    config.updated_by = updated_by
    db.commit()
    db.refresh(config)
    return get_hot_search_config(db)


def import_hot_search_file(db: Session, file_path: str, uploaded_by: str = "operator") -> dict:
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {"imported_count": 0, "failed_rows": [{"row": 1, "reason": "empty_file"}], "categories": []}
    mapping = _headers(rows[0])
    if "category" not in mapping or "keyword" not in mapping:
        return {
            "imported_count": 0,
            "failed_rows": [{"row": 1, "reason": "missing_category_or_keyword_header"}],
            "categories": [],
        }

    batch = HotSearchBatch(filename=Path(file_path).name, uploaded_by=uploaded_by)
    db.add(batch)
    db.flush()
    imported = 0
    failed_rows: list[dict] = []
    categories: dict[str, None] = {}
    for row_no, row in enumerate(rows[1:], start=2):
        category = _clean(_value(row, mapping, "category"))
        keyword = _clean(_value(row, mapping, "keyword"))
        if not category or not keyword:
            failed_rows.append({"row": row_no, "reason": "missing_category_or_keyword"})
            continue
        categories.setdefault(category, None)
        db.add(
            HotSearchTerm(
                batch_id=batch.id,
                category=category,
                keyword=keyword,
                rank=_int(_value(row, mapping, "rank")) or imported + 1,
                search_uv_index=_number(_value(row, mapping, "search_uv_index")),
                search_uv_growth=_number(_value(row, mapping, "search_uv_growth")),
                click_rate_index=_number(_value(row, mapping, "click_rate_index")),
                opportunity_index=_number(_value(row, mapping, "opportunity_index")),
                gmv_index=_number(_value(row, mapping, "gmv_index")),
                sales_index=_number(_value(row, mapping, "sales_index")),
            )
        )
        imported += 1
    db.commit()
    return {
        "batch_id": batch.id,
        "filename": batch.filename,
        "imported_count": imported,
        "failed_rows": failed_rows,
        "categories": list(categories.keys()),
    }


def _avoid_terms_for_product(db: Session, product: Product) -> list[str]:
    line = _product_line(product)
    if not line:
        return []
    rule = (
        db.query(Rule)
        .filter(Rule.rule_type == "hot_search_avoid_terms", Rule.rule_name == line, Rule.enabled.is_(True))
        .first()
    )
    return _split_terms(rule.content if rule else "")


def _has_gender_conflict(keyword: str, gender: str) -> bool:
    if gender == "男":
        return "女" in keyword
    if gender == "女":
        return "男" in keyword
    return False


def select_hot_terms_for_product(db: Session, product: Product, limit: int = 10) -> dict:
    category = _clean(product.category_3)
    if not category:
        return {
            "selected_hot_terms": [],
            "excluded_hot_terms": [],
            "hot_search_source_batch": None,
        }
    batch = db.query(HotSearchBatch).order_by(HotSearchBatch.created_at.desc(), HotSearchBatch.id.desc()).first()
    if not batch:
        return {
            "selected_hot_terms": [],
            "excluded_hot_terms": [],
            "hot_search_source_batch": None,
        }
    terms = (
        db.query(HotSearchTerm)
        .filter(HotSearchTerm.batch_id == batch.id, HotSearchTerm.category == category)
        .order_by(HotSearchTerm.rank.asc(), HotSearchTerm.id.asc())
        .all()
    )
    avoid_terms = _avoid_terms_for_product(db, product)
    gender = _clean(product.gender)
    selected: list[str] = []
    excluded: list[dict] = []
    for term in terms:
        keyword = _clean(term.keyword)
        avoid_term = next((item for item in avoid_terms if item and item in keyword), None)
        if avoid_term:
            excluded.append({"keyword": keyword, "reason": f"avoid_term:{avoid_term}"})
            continue
        if _has_gender_conflict(keyword, gender):
            excluded.append({"keyword": keyword, "reason": "gender_conflict"})
            continue
        if keyword and keyword not in selected:
            selected.append(keyword)
        if len(selected) >= limit:
            break
    return {
        "selected_hot_terms": selected,
        "excluded_hot_terms": excluded,
        "hot_search_source_batch": batch.id,
    }


def evaluate_hot_title(title: str | None, hot_terms: list[str] | None) -> dict:
    title_text = _clean(title)
    terms = [_clean(term) for term in (hot_terms or []) if _clean(term)]
    matched = [term for term in terms if term in title_text]
    missing = [term for term in terms if term not in title_text]
    return {"matched_hot_terms": matched, "missing_hot_terms": missing}
