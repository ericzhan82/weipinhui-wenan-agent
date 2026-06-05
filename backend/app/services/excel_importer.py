import re

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import Product, ProductSku

ALIASES = {
    "style_no": ["款号", "款式", "商品款号", "产品款号", "spu", "SPU"],
    "product_no": ["货号", "商品货号", "产品货号", "唯品货号", "商品编码"],
    "fba": ["FBA", "设计师卖点", "产品卖点", "商品卖点", "核心卖点", "FBA设计师卖点"],
    "category_3": ["三级品类", "三类目", "三级分类", "三级分类名称"],
    "category_4": ["四级品类", "四类目", "四级分类", "四级分类名称"],
    "age_range": ["适用岁段", "适应岁段", "年龄段", "尺码段"],
    "gender": ["性别"],
    "season": ["季节"],
    "scene": ["场景"],
    "color_name": ["颜色", "颜色名称", "色系"],
    "color_code": ["色号"],
    "sku_no": ["SKC", "SKC/货号", "sku", "SKU"],
}
PRODUCT_FIELDS = ("category_3", "category_4", "age_range", "gender", "season", "scene", "fba")


def _normalize_header(value: str) -> str:
    return re.sub(r"[\s:：/\\()（）_\-]+", "", str(value or "")).lower()


def _header_map(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalized = {str(header or "").strip(): index for index, header in enumerate(headers)}
    compact = {_normalize_header(header): index for index, header in enumerate(headers) if str(header or "").strip()}
    for field, names in ALIASES.items():
        for name in names:
            if name in normalized:
                mapping[field] = normalized[name]
                break
            key = _normalize_header(name)
            if key in compact:
                mapping[field] = compact[key]
                break
            for header_key, index in compact.items():
                if key and key in header_key:
                    mapping[field] = index
                    break
            if field in mapping:
                break
    return mapping


def _find_header_row(rows: list[tuple], max_scan_rows: int = 10) -> tuple[int, dict[str, int]]:
    best_row_index = 0
    best_mapping: dict[str, int] = {}
    for row_index, row in enumerate(rows[:max_scan_rows]):
        mapping = _header_map([str(cell or "").strip() for cell in row])
        if "style_no" in mapping or "product_no" in mapping:
            return row_index, mapping
        if len(mapping) > len(best_mapping):
            best_row_index = row_index
            best_mapping = mapping
    return best_row_index, best_mapping


def _value(row: tuple, mapping: dict[str, int], field: str) -> str | None:
    index = mapping.get(field)
    if index is None or index >= len(row):
        return None
    value = row[index]
    return str(value).strip() if value is not None else None


def _is_empty_row(row: tuple) -> bool:
    return all(str(value).strip() == "" for value in row if value is not None)


def _context_ready(product: Product) -> bool:
    return bool(
        (product.category_3 or "").strip()
        and (product.category_4 or "").strip()
        and (product.fba or "").strip()
    )


def _sync_product_status(product: Product) -> None:
    if product.copy_output and product.copy_output.title:
        product.status = "generated"
    elif _context_ready(product):
        product.status = "ready"
    elif product.status not in {"generating", "failed"}:
        product.status = "draft"


def _apply_context_values(product: Product, values: dict[str, str | None], operator_name: str) -> bool:
    changed = False
    for field in PRODUCT_FIELDS:
        incoming = values.get(field)
        if incoming and not (getattr(product, field) or "").strip():
            setattr(product, field, incoming)
            changed = True
    if changed:
        product.updated_by = operator_name
    _sync_product_status(product)
    return changed


def _sku_exists(db: Session, product_id: int, sku_no: str | None, color_name: str | None, color_code: str | None) -> bool:
    query = db.query(ProductSku).filter(ProductSku.product_id == product_id)
    if sku_no:
        query = query.filter(ProductSku.sku_no == sku_no)
    elif color_name:
        query = query.filter(ProductSku.color_name == color_name)
    else:
        return False
    if color_code:
        query = query.filter(ProductSku.color_code == color_code)
    return query.first() is not None


def import_excel(db: Session, file_path: str, workspace_id: int | None = None, operator_name: str = "excel") -> dict:
    workbook = load_workbook(file_path)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {"success_count": 0, "failed_rows": [{"row": 1, "reason": "空文件"}], "skipped_count": 0}
    header_row_index, mapping = _find_header_row(rows)
    success = 0
    skipped = 0
    updated = 0
    failed: list[dict] = []
    for row_no, row in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
        if _is_empty_row(row):
            continue
        style_no = _value(row, mapping, "style_no")
        product_no = _value(row, mapping, "product_no")
        if not style_no and not product_no:
            failed.append({"row": row_no, "reason": "缺少款号或货号"})
            continue
        existing = (
            db.query(Product)
            .filter(Product.style_no == style_no, Product.product_no == product_no)
            .filter(Product.workspace_id == workspace_id)
            .first()
        )
        if existing:
            skipped += 1
            product = existing
        else:
            values = {field: _value(row, mapping, field) for field in PRODUCT_FIELDS}
            product = Product(
                style_no=style_no,
                product_no=product_no,
                **values,
                workspace_id=workspace_id,
                created_by=operator_name,
                updated_by=operator_name,
            )
            _sync_product_status(product)
            db.add(product)
            db.flush()
            success += 1
        if existing:
            values = {field: _value(row, mapping, field) for field in PRODUCT_FIELDS}
            if _apply_context_values(product, values, operator_name):
                updated += 1
        sku_no = _value(row, mapping, "sku_no")
        color_name = _value(row, mapping, "color_name")
        color_code = _value(row, mapping, "color_code")
        if (sku_no or color_name) and not _sku_exists(db, product.id, sku_no, color_name, color_code):
            db.add(
                ProductSku(
                    product_id=product.id,
                    sku_no=sku_no,
                    color_name=color_name,
                    color_code=color_code,
                )
            )
    db.commit()
    return {"success_count": success, "failed_rows": failed, "skipped_count": skipped, "updated_count": updated}
